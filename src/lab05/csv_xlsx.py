"""Модуль конвертации CSV → XLSX (ЛР5)."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str | Path, xlsx_path: str | Path) -> None:
    """Конвертирует CSV в XLSX с автошириной колонок."""
    p_in = Path(csv_path)
    p_out = Path(xlsx_path)
    
    with p_in.open("r", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)
        
    if not rows: raise ValueError("Пустой CSV")
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    col_widths = {}
    for row_idx, row in enumerate(rows, start=1):
        ws.append(row)
        for col_idx, value in enumerate(row, start=1):
            length = max(len(str(value)), 8) # не менее 8 символов
            if col_idx not in col_widths or length > col_widths[col_idx]: col_widths[col_idx] = length
                
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
        
    p_out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p_out)